from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import PieChart, BarChart, RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from datetime import date

OUT = "templates/Application-Tracker-Pro-Sample.xlsx"

NAVY = "12213D"; TEAL = "0E7C66"; CYAN = "35C6E8"; LIGHT = "F4F6FA"; MID = "E1F5EE"; GRAY = "5F5E5A"; WHITE = "FFFFFF"
FONT = "Arial"

def f(size=10, bold=False, color="000000", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin = Side(style="thin", color="D0D4DC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", fgColor=c)

wb = Workbook()
wb.properties.creator = "ApplyAbroadLab.com"
wb.properties.title = "Apply Abroad Lab Application Tracker Pro"

def banner(ws, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(row=1, column=1, value="ApplyAbroadLab.com — Application Tracker Pro")
    c.hyperlink = "https://applyabroadlab.com/"
    c.font = f(11, True, WHITE, underline="single"); c.fill = fill(NAVY); c.alignment = CENTER
    ws.row_dimensions[1].height = 24

def header_row(ws, row, headers, widths, left_cols=()):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f(10, True, WHITE); c.fill = fill(NAVY); c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

# ------------------------------------------------------------------ Lists
lists = wb.active; lists.title = "Lists"
STATUS = ["Not Started", "In Progress", "Submitted", "Interview", "Accepted", "Rejected"]
PRIORITY = ["High", "Medium", "Low"]
LEVEL = ["PhD", "Postdoc", "Master"]
FIELDS = ["Biology", "Cancer Biology", "Chemistry", "Computer Science", "Psychology", "Civil Engineering",
          "Mechanical Engineering", "Architecture", "Business & Management", "Marketing", "Food Biotechnology",
          "Computational Biology", "Molecular Oncology"]
CONTACT = ["Not contacted", "Awaiting reply", "Replied", "Interview scheduled"]
for col, (name, vals) in enumerate({"Status": STATUS, "Priority": PRIORITY, "Level": LEVEL, "Field": FIELDS,
                                     "Contact Status": CONTACT}.items(), start=1):
    lists.cell(row=1, column=col, value=name).font = f(10, True)
    for r, v in enumerate(vals, start=2):
        lists.cell(row=r, column=col, value=v)
    lists.column_dimensions[get_column_letter(col)].width = 24

# ------------------------------------------------------------------ Universities & Supervisors
uni = wb.create_sheet("Universities & Supervisors")
UNIS = [
 ("Technical University of Munich","Germany","Munich","Computational Biology","Prof. Lena Hartmann",9,7,"Fully funded (DFG)","Replied","https://www.tum.de"),
 ("Heidelberg University","Germany","Heidelberg","Molecular Oncology","Prof. Jonas Weber",9,5,"Fully funded (DKFZ)","Awaiting reply","https://www.uni-heidelberg.de"),
 ("RWTH Aachen University","Germany","Aachen","Mechanical Engineering","Prof. Katrin Vogel",8,6,"Fully funded (Institute)","Not contacted","https://www.rwth-aachen.de"),
 ("Delft University of Technology","Netherlands","Delft","Civil Engineering","Dr. Pieter de Vries",9,8,"Fully funded (University)","Replied","https://www.tudelft.nl"),
 ("University of Amsterdam","Netherlands","Amsterdam","Psychology","Dr. Sanne Bakker",8,7,"Fully funded (NWO)","Interview scheduled","https://www.uva.nl"),
 ("Wageningen University","Netherlands","Wageningen","Food Biotechnology","Prof. Joris van Dijk",8,9,"Fully funded (University)","Replied","https://www.wur.nl"),
 ("Utrecht University","Netherlands","Utrecht","Biology","Dr. Femke Jansen",8,6,"Fully funded (University)","Awaiting reply","https://www.uu.nl"),
 ("Karolinska Institutet","Sweden","Stockholm","Cancer Biology","Prof. Anna Lindqvist",10,6,"Fully funded (KI)","Replied","https://ki.se"),
 ("KTH Royal Institute of Technology","Sweden","Stockholm","Computer Science","Dr. Erik Johansson",8,8,"Fully funded (WASP)","Awaiting reply","https://www.kth.se"),
 ("Lund University","Sweden","Lund","Chemistry","Prof. Karin Svensson",8,7,"Fully funded (University)","Replied","https://www.lu.se"),
 ("Politecnico di Milano","Italy","Milan","Architecture","Prof. Giulia Ferrari",8,5,"Scholarship (Regional)","Not contacted","https://www.polimi.it"),
 ("University of Bologna","Italy","Bologna","Business & Management","Prof. Marco Rossi",7,6,"Scholarship (University)","Replied","https://www.unibo.it"),
 ("KU Leuven","Belgium","Leuven","Chemistry","Prof. Bram Peeters",9,7,"Fully funded (FWO)","Awaiting reply","https://www.kuleuven.be"),
 ("Ghent University","Belgium","Ghent","Biology","Dr. Elise Claes",8,8,"Fully funded (BOF)","Replied","https://www.ugent.be"),
 ("ETH Zurich","Switzerland","Zurich","Computer Science","Prof. Daniel Meier",10,4,"Fully funded (ETH)","Not contacted","https://ethz.ch"),
 ("EPFL","Switzerland","Lausanne","Mechanical Engineering","Dr. Claire Dubois",10,6,"Fully funded (SNSF)","Awaiting reply","https://www.epfl.ch"),
 ("NTNU","Norway","Trondheim","Civil Engineering","Prof. Ola Nilsen",8,9,"Fully funded (University)","Replied","https://www.ntnu.edu"),
 ("University of Oslo","Norway","Oslo","Psychology","Dr. Ingrid Haugen",8,7,"Fully funded (University)","Interview scheduled","https://www.uio.no"),
 ("University of Toronto","Canada","Toronto","Marketing","Prof. Sarah Mitchell",9,5,"Partially funded","Not contacted","https://www.utoronto.ca"),
 ("McGill University","Canada","Montreal","Cancer Biology","Dr. Olivier Tremblay",9,6,"Fully funded (CIHR)","Replied","https://www.mcgill.ca"),
]
banner(uni, 10)
uni.cell(row=2, column=1, value="Reference table — one row per university / supervisor. Applications pull Country, Supervisor and Funding from here automatically.").font = f(9, italic=True, color=GRAY)
uni.merge_cells("A2:J2")
header_row(uni, 3, ["University","Country","City","Research Area","Supervisor","Prestige (1-10)","Responsiveness (1-10)","Funding Type","Contact Status","Website"],
           [30,13,12,24,22,13,15,24,18,28])
for r, row in enumerate(UNIS, start=4):
    for c, v in enumerate(row, start=1):
        cell = uni.cell(row=r, column=c, value=v)
        cell.font = f(); cell.border = BORDER
        cell.alignment = LEFT if c in (4, 10) else CENTER
        if c == 10:
            cell.hyperlink = v; cell.font = f(10, color="1565C0", underline="single")
    uni.row_dimensions[r].height = 20
U_FIRST, U_LAST = 4, 3 + len(UNIS)
dv = DataValidation(type="list", formula1=f"=Lists!$E$2:$E${1+len(CONTACT)}", allow_blank=True)
uni.add_data_validation(dv); dv.add(f"I{U_FIRST}:I{U_LAST}")
uni.freeze_panes = "B4"

# ------------------------------------------------------------------ Applications
app = wb.create_sheet("Applications")
APPS = [
 # University, Program, Field, Level, Deadline, Status, Priority, Fit, Docs%, LastContact, Notes
 ("Technical University of Munich","PhD in Computational Biology","Computational Biology","PhD",date(2026,10,5),"In Progress","High",9,80,date(2026,9,15),"Supervisor confirmed interest; finalise research proposal"),
 ("Heidelberg University","PhD in Molecular Oncology (DKFZ)","Molecular Oncology","PhD",date(2026,11,15),"In Progress","High",8,60,date(2026,9,10),"Two LORs received, third pending"),
 ("RWTH Aachen University","PhD in Mechanical Engineering","Mechanical Engineering","PhD",date(2027,1,31),"Not Started","Medium",7,20,None,"Need to email supervisor before drafting SOP"),
 ("Delft University of Technology","PhD in Structural Engineering","Civil Engineering","PhD",date(2026,10,12),"Submitted","High",9,100,date(2026,9,2),"Submitted via portal; acknowledgement received"),
 ("University of Amsterdam","PhD in Clinical Psychology","Psychology","PhD",date(2026,10,18),"Interview","High",8,100,date(2026,9,18),"Interview scheduled 28 Sep, prepare 10-min presentation"),
 ("Wageningen University","PhD in Food Biotechnology","Food Biotechnology","PhD",date(2026,12,1),"In Progress","Medium",8,70,date(2026,9,12),"Supervisor asked for detailed CV"),
 ("Utrecht University","PhD in Developmental Biology","Biology","PhD",date(2026,12,15),"Not Started","Low",6,10,None,"Backup option"),
 ("Karolinska Institutet","Postdoc in Cancer Biology","Cancer Biology","Postdoc",date(2026,10,20),"Submitted","High",9,100,date(2026,9,5),"Strong fit; awaiting shortlist"),
 ("KTH Royal Institute of Technology","PhD in Machine Learning (WASP)","Computer Science","PhD",date(2026,11,30),"In Progress","High",8,50,date(2026,9,8),"Coding sample required"),
 ("Lund University","PhD in Analytical Chemistry","Chemistry","PhD",date(2027,1,15),"In Progress","Medium",7,40,date(2026,9,1),"Draft SOP version 2 in progress"),
 ("Politecnico di Milano","PhD in Architecture & Urban Design","Architecture","PhD",date(2027,2,28),"Not Started","Low",6,0,None,"Check scholarship eligibility"),
 ("University of Bologna","PhD in Management","Business & Management","PhD",date(2027,3,15),"In Progress","Medium",7,30,date(2026,8,28),"Portfolio of publications needed"),
 ("KU Leuven","PhD in Organic Chemistry","Chemistry","PhD",date(2026,11,10),"In Progress","High",8,65,date(2026,9,14),"FWO fellowship co-application"),
 ("Ghent University","PhD in Plant Biology","Biology","PhD",date(2026,12,20),"Submitted","Medium",7,100,date(2026,9,3),"Submitted; reference letters uploaded"),
 ("ETH Zurich","PhD in Computer Science","Computer Science","PhD",date(2026,12,10),"Not Started","Medium",9,15,None,"Very competitive; prioritise supervisor outreach"),
 ("EPFL","Postdoc in Robotics","Mechanical Engineering","Postdoc",date(2027,1,20),"In Progress","Medium",8,45,date(2026,9,9),"Waiting on supervisor reply about SNSF budget"),
 ("NTNU","PhD in Geotechnical Engineering","Civil Engineering","PhD",date(2026,11,25),"Interview","High",8,100,date(2026,9,17),"Interview 30 Sep, technical questions expected"),
 ("University of Oslo","PhD in Cognitive Psychology","Psychology","PhD",date(2027,2,10),"Interview","Medium",7,100,date(2026,9,16),"Second-round interview"),
 ("University of Toronto","PhD in Marketing","Marketing","PhD",date(2027,3,1),"Rejected","Low",6,100,date(2026,8,20),"Rejected on funding grounds; keep contact"),
 ("McGill University","Postdoc in Cancer Biology","Cancer Biology","Postdoc",date(2026,12,5),"Accepted","Low",9,100,date(2026,9,19),"Offer received; decision pending"),
]
banner(app, 16)
app.cell(row=2, column=1, value="Main tracker — edit the white columns. Grey columns (Country, Supervisor, Funding, Days Until Deadline) are calculated automatically.").font = f(9, italic=True, color=GRAY)
app.merge_cells("A2:P2")
H = ["ID","University","Country","Program","Field","Level","Supervisor","Deadline","Days Until Deadline","Status","Priority",
     "Research Fit (1-10)","Funding Type","Documents Ready (%)","Last Contact","Notes","Upcoming Key","Match Index"]
W = [6,30,13,30,20,10,22,13,12,13,10,11,24,12,13,40,10,10]
header_row(app, 3, H, W)
A_FIRST = 4; A_LAST = 3 + len(APPS)
UR = f"'Universities & Supervisors'!$A${U_FIRST}:$A${U_LAST}"
def ucol(col): return f"'Universities & Supervisors'!${col}${U_FIRST}:${col}${U_LAST}"
for i, row in enumerate(APPS):
    r = A_FIRST + i
    uname, prog, field, level, dl, status, prio, fit, docs, last, notes = row
    vals = {1: i+1, 2: uname, 3: f'=IFERROR(INDEX({ucol("B")},MATCH(B{r},{UR},0)),"")', 4: prog, 5: field, 6: level,
            7: f'=IFERROR(INDEX({ucol("E")},MATCH(B{r},{UR},0)),"")', 8: dl, 9: f'=IF(H{r}="","",H{r}-TODAY())',
            10: status, 11: prio, 12: fit, 13: f'=IFERROR(INDEX({ucol("H")},MATCH(B{r},{UR},0)),"")', 14: docs/100,
            15: last, 16: notes,
            17: f'=IF(AND(I{r}<>"",I{r}>=0,I{r}<=30),I{r}+ROW()/100000,"")',
            18: f'=IF(AND(OR(Explorer!$C$4="All",C{r}=Explorer!$C$4),OR(Explorer!$C$5="All",J{r}=Explorer!$C$5),OR(Explorer!$C$6="All",K{r}=Explorer!$C$6)),MAX($R${A_FIRST-1}:R{r-1})+1,"")'}
    for c, v in vals.items():
        cell = app.cell(row=r, column=c, value=v)
        cell.font = f(); cell.border = BORDER
        cell.alignment = LEFT if c in (4, 16) else CENTER
        if c in (3, 7, 9, 13): cell.fill = fill(LIGHT)
        if c == 8 or c == 15: cell.number_format = "dd-mmm-yyyy"
        if c == 14: cell.number_format = "0%"
    app.row_dimensions[r].height = 20
# Validations
def add_dv(ws, rng, formula):
    d = DataValidation(type="list", formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(rng)
add_dv(app, f"B{A_FIRST}:B{A_LAST+30}", f"={UR}")
add_dv(app, f"E{A_FIRST}:E{A_LAST+30}", f"=Lists!$D$2:$D${1+len(FIELDS)}")
add_dv(app, f"F{A_FIRST}:F{A_LAST+30}", f"=Lists!$C$2:$C${1+len(LEVEL)}")
add_dv(app, f"J{A_FIRST}:J{A_LAST+30}", f"=Lists!$A$2:$A${1+len(STATUS)}")
add_dv(app, f"K{A_FIRST}:K{A_LAST+30}", f"=Lists!$B$2:$B${1+len(PRIORITY)}")
dvn = DataValidation(type="whole", operator="between", formula1="1", formula2="10"); app.add_data_validation(dvn); dvn.add(f"L{A_FIRST}:L{A_LAST+30}")
# Conditional formatting
rngI = f"I{A_FIRST}:I{A_LAST+30}"
app.conditional_formatting.add(rngI, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9E9E9E"), fill=fill("EEEEEE")))
app.conditional_formatting.add(rngI, CellIsRule(operator="between", formula=["0","30"], font=Font(color="C62828", bold=True), fill=fill("FDECEA")))
app.conditional_formatting.add(rngI, CellIsRule(operator="between", formula=["31","60"], font=Font(color="E68A00", bold=True), fill=fill("FFF4E0")))
app.conditional_formatting.add(rngI, CellIsRule(operator="greaterThan", formula=["60"], font=Font(color="2E7D32"), fill=fill("EAF3DE")))
rngJ = f"J{A_FIRST}:J{A_LAST+30}"
for s, col in {"Accepted":"C8E6C9","Rejected":"FFCDD2","Interview":"BBDEFB","Submitted":"E1F5EE","Not Started":"EEEEEE"}.items():
    app.conditional_formatting.add(rngJ, FormulaRule(formula=[f'$J{A_FIRST}="{s}"'], fill=fill(col)))
rngK = f"K{A_FIRST}:K{A_LAST+30}"
app.conditional_formatting.add(rngK, FormulaRule(formula=[f'$K{A_FIRST}="High"'], font=Font(color="C62828", bold=True)))
app.conditional_formatting.add(f"N{A_FIRST}:N{A_LAST+30}", CellIsRule(operator="equal", formula=["1"], font=Font(color="2E7D32", bold=True)))
app.column_dimensions["Q"].hidden = True; app.column_dimensions["R"].hidden = True
app.freeze_panes = "C4"
app.auto_filter.ref = f"A3:P{A_LAST}"

# ------------------------------------------------------------------ Scoring (radar)
sc = wb.create_sheet("Scoring")
banner(sc, 8)
sc.cell(row=2, column=1, value="Shortlist comparison — pick up to 4 universities in the yellow cells. All axes are normalised to 0-10 so the radar shape is comparable.").font = f(9, italic=True, color=GRAY)
sc.merge_cells("A2:H2")
header_row(sc, 3, ["Axis","Shortlist 1","Shortlist 2","Shortlist 3","Shortlist 4","","Scoring rule",""], [26,26,26,26,26,3,60,3])
sc.merge_cells("G3:H3")
picks = ["Technical University of Munich","Karolinska Institutet","University of Bologna","ETH Zurich"]
sc.cell(row=4, column=1, value="University").font = f(10, True)
sc.cell(row=4, column=1).alignment = CENTER; sc.cell(row=4, column=1).border = BORDER
for j, p in enumerate(picks, start=2):
    c = sc.cell(row=4, column=j, value=p); c.fill = fill("FFF9C4"); c.font = f(10, True); c.alignment = CENTER; c.border = BORDER
add_dv(sc, "B4:E4", f"={UR}")
AR = f"Applications!$B${A_FIRST}:$B${A_LAST}"
def acol(col): return f"Applications!${col}${A_FIRST}:${col}${A_LAST}"
axes = [
 ("Research Fit", lambda L: f'=IFERROR(INDEX({acol("L")},MATCH({L}$4,{AR},0)),0)', "Research Fit column of the application, already on a 1-10 scale"),
 ("Funding Strength", lambda L: f'=IFERROR(IF(ISNUMBER(SEARCH("Fully",INDEX({acol("M")},MATCH({L}$4,{AR},0)))),10,IF(ISNUMBER(SEARCH("Partial",INDEX({acol("M")},MATCH({L}$4,{AR},0)))),6,7)),0)', "Fully funded = 10, Scholarship = 7, Partially funded = 6"),
 ("Prestige", lambda L: f'=IFERROR(INDEX({ucol("F")},MATCH({L}$4,{UR},0)),0)', "Prestige score from the Universities sheet (1-10)"),
 ("Deadline Urgency", lambda L: f'=IFERROR(MAX(0,MIN(10,ROUND(10-INDEX({acol("I")},MATCH({L}$4,{AR},0))/20,0))),0)', "10 = deadline now, minus 1 point per 20 days away, floor 0"),
 ("Supervisor Responsiveness", lambda L: f'=IFERROR(INDEX({ucol("G")},MATCH({L}$4,{UR},0)),0)', "Responsiveness score from the Universities sheet (1-10)"),
]
for i, (name, fn, rule) in enumerate(axes, start=5):
    a = sc.cell(row=i, column=1, value=name); a.font = f(10, True); a.alignment = CENTER; a.border = BORDER; a.fill = fill(MID)
    for j, L in enumerate("BCDE", start=2):
        c = sc.cell(row=i, column=j, value=fn(L)); c.font = f(); c.alignment = CENTER; c.border = BORDER
    g = sc.cell(row=i, column=7, value=rule); g.font = f(9, color=GRAY); g.alignment = LEFT; sc.merge_cells(start_row=i, start_column=7, end_row=i, end_column=8)
    sc.row_dimensions[i].height = 20
t = sc.cell(row=10, column=1, value="Overall (average)"); t.font = f(10, True); t.alignment = CENTER; t.border = BORDER; t.fill = fill(NAVY); t.font = f(10, True, WHITE)
for j, L in enumerate("BCDE", start=2):
    c = sc.cell(row=10, column=j, value=f"=ROUND(AVERAGE({L}5:{L}9),1)"); c.font = f(10, True); c.alignment = CENTER; c.border = BORDER; c.fill = fill(MID)
radar = RadarChart(); radar.type = "filled"; radar.style = 26; radar.title = "Shortlist comparison (0-10)"
radar.add_data(Reference(sc, min_col=2, max_col=5, min_row=4, max_row=9), titles_from_data=True)
radar.set_categories(Reference(sc, min_col=1, min_row=5, max_row=9))
radar.y_axis.scaling.min = 0; radar.y_axis.scaling.max = 10
radar.height = 10; radar.width = 18
sc.add_chart(radar, "A13")

# ------------------------------------------------------------------ Explorer
ex = wb.create_sheet("Explorer")
banner(ex, 8)
ex.cell(row=2, column=1, value="Interactive filter — change the three yellow selectors and the list, KPIs and chart update instantly.").font = f(9, italic=True, color=GRAY)
ex.merge_cells("A2:H2")
for r, (lab, key) in enumerate([("Country","B"), ("Status","A"), ("Priority","B")], start=4):
    l = ex.cell(row=r, column=2, value=lab); l.font = f(10, True); l.alignment = CENTER; l.border = BORDER; l.fill = fill(MID)
    v = ex.cell(row=r, column=3, value="All"); v.font = f(10, True); v.alignment = CENTER; v.border = BORDER; v.fill = fill("FFF9C4")
# selector lists on Lists sheet (with "All")
lists.cell(row=1, column=7, value="Country+All").font = f(10, True)
countries = ["All"] + sorted({u[1] for u in UNIS})
for r, v in enumerate(countries, start=2): lists.cell(row=r, column=7, value=v)
lists.cell(row=1, column=8, value="Status+All").font = f(10, True)
for r, v in enumerate(["All"]+STATUS, start=2): lists.cell(row=r, column=8, value=v)
lists.cell(row=1, column=9, value="Priority+All").font = f(10, True)
for r, v in enumerate(["All"]+PRIORITY, start=2): lists.cell(row=r, column=9, value=v)
add_dv(ex, "C4", f"=Lists!$G$2:$G${1+len(countries)}")
add_dv(ex, "C5", f"=Lists!$H$2:$H${2+len(STATUS)}")
add_dv(ex, "C6", f"=Lists!$I$2:$I${2+len(PRIORITY)}")
# KPIs for the selection
kp = [("Matching applications", f"=COUNT(Applications!R{A_FIRST}:R{A_LAST})"),
      ("Average research fit", f'=IFERROR(ROUND(SUMPRODUCT((Applications!R{A_FIRST}:R{A_LAST}<>"")*Applications!L{A_FIRST}:L{A_LAST})/$F$4,1),0)'),
      ("Average documents ready", f'=IFERROR(SUMPRODUCT((Applications!R{A_FIRST}:R{A_LAST}<>"")*Applications!N{A_FIRST}:N{A_LAST})/$F$4,0)'),
      ("Nearest deadline (days)", f'=IFERROR(SUMPRODUCT(MIN(IF(Applications!R{A_FIRST}:R{A_LAST}<>"",IF(Applications!I{A_FIRST}:I{A_LAST}>=0,Applications!I{A_FIRST}:I{A_LAST})))),"")')]
for r, (lab, fo) in enumerate(kp, start=4):
    l = ex.cell(row=r, column=5, value=lab); l.font = f(10, True); l.alignment = CENTER; l.border = BORDER; l.fill = fill(MID)
    v = ex.cell(row=r, column=6, value=fo); v.font = f(12, True, NAVY); v.alignment = CENTER; v.border = BORDER
    if r == 6: v.number_format = "0%"
    ex.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
for r in range(4, 8): ex.row_dimensions[r].height = 22
header_row(ex, 9, ["#","University","Country","Program","Deadline","Days","Status","Priority"], [5,30,13,32,13,8,13,10])
for n in range(1, 21):
    r = 9 + n
    ex.cell(row=r, column=1, value=n)
    src = {2:"B",3:"C",4:"D",5:"H",6:"I",7:"J",8:"K"}
    for c, col in src.items():
        ex.cell(row=r, column=c, value=f'=IFERROR(INDEX(Applications!${col}${A_FIRST}:${col}${A_LAST},MATCH($A{r},Applications!$R${A_FIRST}:$R${A_LAST},0)),"")')
    for c in range(1, 9):
        cell = ex.cell(row=r, column=c); cell.font = f(); cell.border = BORDER
        cell.alignment = LEFT if c == 4 else CENTER
    ex.cell(row=r, column=5).number_format = "dd-mmm-yyyy"
    ex.row_dimensions[r].height = 18
ex.conditional_formatting.add("F10:F29", CellIsRule(operator="between", formula=["0","30"], font=Font(color="C62828", bold=True)))
ex.conditional_formatting.add("F10:F29", CellIsRule(operator="between", formula=["31","60"], font=Font(color="E68A00", bold=True)))
ex.conditional_formatting.add("A10:H29", FormulaRule(formula=['$B10=""'], font=Font(color="FFFFFF"), border=Border()))
# status counts for selection -> bar chart
ex.cell(row=31, column=2, value="Status (selection)").font = f(10, True); ex.cell(row=31, column=3, value="Count").font = f(10, True)
for i, s in enumerate(STATUS, start=32):
    ex.cell(row=i, column=2, value=s).font = f()
    ex.cell(row=i, column=3, value=f'=SUMPRODUCT((Applications!$R${A_FIRST}:$R${A_LAST}<>"")*(Applications!$J${A_FIRST}:$J${A_LAST}="{s}"))').font = f()
    ex.cell(row=i, column=2).alignment = CENTER; ex.cell(row=i, column=3).alignment = CENTER
bar2 = BarChart(); bar2.type = "bar"; bar2.style = 10; bar2.title = "Status of selected applications"; bar2.legend = None
bar2.add_data(Reference(ex, min_col=3, min_row=31, max_row=31+len(STATUS)), titles_from_data=True)
bar2.set_categories(Reference(ex, min_col=2, min_row=32, max_row=31+len(STATUS)))
bar2.height = 7; bar2.width = 14
bar2.series[0].graphicalProperties.solidFill = TEAL
ex.add_chart(bar2, "E31")
ex.freeze_panes = "A10"

# ------------------------------------------------------------------ Dashboard
db = wb.create_sheet("Dashboard", 0)
banner(db, 13)
for col in range(1, 14): db.column_dimensions[get_column_letter(col)].width = 13
db.column_dimensions["A"].width = 3
db.merge_cells("B2:M2"); t = db["B2"]; t.value = "Apply Abroad Lab — Application Tracker Pro"
t.font = f(20, True, WHITE); t.fill = fill(NAVY); t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
db.row_dimensions[2].height = 40
db.merge_cells("B3:M3"); s = db["B3"]
s.value = '="Command center for a graduate application campaign  |  Last updated: "&TEXT(TODAY(),"dd-mmm-yyyy")'
s.font = f(10, italic=True, color=GRAY); s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
AI = f"Applications!$I${A_FIRST}:$I${A_LAST}"; AJ = f"Applications!$J${A_FIRST}:$J${A_LAST}"; AK = f"Applications!$K${A_FIRST}:$K${A_LAST}"
AL = f"Applications!$L${A_FIRST}:$L${A_LAST}"; AN = f"Applications!$N${A_FIRST}:$N${A_LAST}"; AB = f"Applications!$B${A_FIRST}:$B${A_LAST}"
kpis = [("TOTAL APPLICATIONS", f'=COUNTA({AB})', None),
        ("NEAREST DEADLINE (DAYS)", f'=SUMPRODUCT(MIN(IF({AI}>=0,{AI})))', None),
        ("HIGH PRIORITY OPEN", f'=SUMPRODUCT(({AK}="High")*({AJ}<>"Submitted")*({AJ}<>"Accepted")*({AJ}<>"Rejected"))', None),
        ("COMPLETION RATE", f'=IFERROR(AVERAGE({AN}),0)', "0%"),
        ("ACCEPTED", f'=COUNTIF({AJ},"Accepted")', None),
        ("AVG. RESEARCH FIT", f'=ROUND(AVERAGE({AL}),1)', None)]
for i, (lab, fo, nf) in enumerate(kpis):
    c1 = 2 + i*2; c2 = c1 + 1
    db.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
    db.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
    l = db.cell(row=5, column=c1, value=lab); l.font = f(9, True, GRAY); l.alignment = CENTER; l.fill = fill(LIGHT)
    v = db.cell(row=6, column=c1, value=fo); v.font = f(24, True, NAVY); v.alignment = CENTER; v.fill = fill(LIGHT)
    if nf: v.number_format = nf
    for cc in (c1, c2):
        db.cell(row=5, column=cc).border = BORDER; db.cell(row=6, column=cc).border = BORDER
db.row_dimensions[5].height = 22; db.row_dimensions[6].height = 46
db["D6"].font = f(24, True, TEAL)
# chart data helper block (columns O:R)
db["O4"] = "Chart data (auto)"; db["O4"].font = f(9, True, GRAY)
db["O5"] = "Status"; db["P5"] = "Count"
for i, st in enumerate(STATUS, start=6):
    db.cell(row=i, column=15, value=st); db.cell(row=i, column=16, value=f'=COUNTIF({AJ},"{st}")')
db["O13"] = "Country"; db["P13"] = "Count"
AC = f"Applications!$C${A_FIRST}:$C${A_LAST}"
for i, co in enumerate(countries[1:], start=14):
    db.cell(row=i, column=15, value=co); db.cell(row=i, column=16, value=f'=COUNTIF({AC},"{co}")')
for r in range(5, 14+len(countries)):
    for c in (15, 16):
        db.cell(row=r, column=c).font = f(9, color=GRAY); db.cell(row=r, column=c).alignment = CENTER
db.column_dimensions["O"].width = 14; db.column_dimensions["P"].width = 8
# section header
for col, title in [(2, "APPLICATION STATUS"), (6, "APPLICATIONS BY COUNTRY"), (10, "SHORTLIST COMPARISON")]:
    db.merge_cells(start_row=8, start_column=col, end_row=8, end_column=col+3)
    h = db.cell(row=8, column=col, value=title); h.font = f(10, True, NAVY); h.alignment = Alignment(horizontal="left", vertical="center")
db.row_dimensions[8].height = 20
pie = PieChart(); pie.title = None; pie.style = 26
pie.add_data(Reference(db, min_col=16, min_row=5, max_row=5+len(STATUS)), titles_from_data=True)
pie.set_categories(Reference(db, min_col=15, min_row=6, max_row=5+len(STATUS)))
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True; pie.dataLabels.showVal = False; pie.dataLabels.showCatName = True; pie.dataLabels.showSerName = False; pie.dataLabels.showLeaderLines = True; pie.legend.position = "b"
pcols = ["B4B2A9", "0E7C66", "35C6E8", "378ADD", "639922", "D85A30"]
for i, col in enumerate(pcols):
    pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = col; pie.series[0].dPt.append(pt)
pie.height = 8.5; pie.width = 12.5
db.add_chart(pie, "B9")
bar = BarChart(); bar.type = "bar"; bar.style = 10; bar.title = None; bar.legend = None
bar.add_data(Reference(db, min_col=16, min_row=13, max_row=13+len(countries)-1), titles_from_data=True)
bar.set_categories(Reference(db, min_col=15, min_row=14, max_row=13+len(countries)-1))
bar.series[0].graphicalProperties.solidFill = NAVY
bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True; bar.dataLabels.showSerName = False; bar.dataLabels.showCatName = False
bar.height = 8.5; bar.width = 12.5; bar.y_axis.majorGridlines = None
db.add_chart(bar, "F9")
radar2 = RadarChart(); radar2.type = "filled"; radar2.style = 26; radar2.title = None
radar2.add_data(Reference(sc, min_col=2, max_col=5, min_row=4, max_row=9), titles_from_data=True)
radar2.set_categories(Reference(sc, min_col=1, min_row=5, max_row=9))
radar2.y_axis.scaling.min = 0; radar2.y_axis.scaling.max = 10
radar2.height = 8.5; radar2.width = 12.5
db.add_chart(radar2, "J9")
# upcoming deadlines table
db.merge_cells("B27:G27"); h = db["B27"]; h.value = "UPCOMING DEADLINES (NEXT 30 DAYS)"; h.font = f(10, True, NAVY)
db.merge_cells("I27:M27"); h2 = db["I27"]; h2.value = "QUICK NAVIGATION"; h2.font = f(10, True, NAVY)
hdr = ["#", "University", "", "Program", "", "Deadline", "Days"]
for c, hv in enumerate(hdr, start=2):
    cell = db.cell(row=28, column=c, value=hv); cell.font = f(9, True, WHITE); cell.fill = fill(NAVY); cell.alignment = CENTER
db.merge_cells("C28:D28"); db.merge_cells("E28:F28")
AQ = f"Applications!$Q${A_FIRST}:$Q${A_LAST}"
for n in range(1, 7):
    r = 28 + n
    db.cell(row=r, column=2, value=n)
    key = f'SMALL({AQ},{n})'
    db.cell(row=r, column=3, value=f'=IFERROR(INDEX({AB},MATCH({key},{AQ},0)),"")')
    db.cell(row=r, column=5, value=f'=IFERROR(INDEX(Applications!$D${A_FIRST}:$D${A_LAST},MATCH({key},{AQ},0)),"")')
    db.cell(row=r, column=7, value=f'=IFERROR(INDEX(Applications!$H${A_FIRST}:$H${A_LAST},MATCH({key},{AQ},0)),"")')
    db.cell(row=r, column=8, value=f'=IFERROR(INDEX({AI},MATCH({key},{AQ},0)),"")')
    db.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4); db.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    for c in range(2, 9):
        cell = db.cell(row=r, column=c); cell.font = f(9); cell.border = BORDER
        cell.alignment = LEFT if c in (3, 5) else CENTER
    db.cell(row=r, column=7).number_format = "dd-mmm-yyyy"
    db.row_dimensions[r].height = 18
db.conditional_formatting.add("H29:H34", CellIsRule(operator="between", formula=["0","30"], font=Font(color="C62828", bold=True)))
db.conditional_formatting.add("B29:H34", FormulaRule(formula=['$C29=""'], font=Font(color="FFFFFF"), border=Border()))
db.cell(row=28, column=8).value = "Days"
# navigation
nav = [("Applications", "Applications"), ("Universities & Supervisors", "Universities & Supervisors"),
       ("Explorer (interactive filter)", "Explorer"), ("Scoring (radar inputs)", "Scoring"), ("How It Works", "How It Works")]
for i, (lab, sheet) in enumerate(nav):
    r = 29 + i
    db.merge_cells(start_row=r, start_column=9, end_row=r, end_column=13)
    c = db.cell(row=r, column=9, value="▶  " + lab)
    c.hyperlink = f"#'{sheet}'!A1"; c.font = f(10, True, WHITE, underline="single"); c.fill = fill(NAVY); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    db.row_dimensions[r].height = 18
# footer
db.merge_cells("B36:M36"); db["B36"] = "Tip: update Status, Priority and Deadline in the Applications sheet — every KPI, chart and list on this page recalculates automatically."
db["B36"].font = f(9, italic=True, color=GRAY)
db.merge_cells("B38:M38"); db["B38"] = "CONNECT WITH APPLY ABROAD LAB"; db["B38"].font = f(10, True, NAVY)
foot = [("Website:", "https://applyabroadlab.com/", "https://applyabroadlab.com/"),
        ("LinkedIn:", "Apply Abroad Lab", "https://www.linkedin.com/in/phd-research-plan-applyabroad/"),
        ("Support:", "via website or LinkedIn", None)]
for i, (lab, txt, link) in enumerate(foot):
    r = 39 + i
    db.cell(row=r, column=2, value=lab).font = f(9, True, NAVY)
    db.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    c = db.cell(row=r, column=3, value=txt); c.font = f(9, color=("1565C0" if link else "000000"), underline=("single" if link else None))
    c.alignment = Alignment(horizontal="left", vertical="center")
    if link: c.hyperlink = link
db.sheet_view.showGridLines = False
db.sheet_view.zoomScale = 100

# ------------------------------------------------------------------ How It Works
hw = wb.create_sheet("How It Works")
banner(hw, 6)
hw.column_dimensions["A"].width = 3; hw.column_dimensions["B"].width = 26; hw.column_dimensions["C"].width = 90
rows = [
 ("What this is", "A single-file application tracker for Master's, PhD and Postdoc campaigns. Everything on the Dashboard is calculated live from the Applications sheet."),
 ("Data model", "Two tables. Applications = one row per application (the fact table). Universities & Supervisors = one row per institution (the reference table). Applications pull Country, Supervisor and Funding from the reference table by University name, so a fact is entered once and reused everywhere."),
 ("What is automated", "Days-until-deadline countdown with red/orange/green alerts, six live KPI cards, status pie chart, country bar chart, 4-way shortlist radar chart, upcoming-deadline list (next 30 days), interactive Explorer filter, and the shortlist scoring table."),
 ("Adding an application", "Go to Applications, type in the next empty row: pick University from the dropdown (Country, Supervisor and Funding fill in automatically), enter Program, Field, Level, Deadline, Status, Priority, Research Fit and Documents Ready. Takes about 60 seconds."),
 ("Adding a university", "Add a row to Universities & Supervisors with its Prestige and Responsiveness scores. It becomes selectable in Applications and Scoring immediately."),
 ("Comparing a shortlist", "Open Scoring, pick up to four universities in the yellow cells. Every axis is normalised to 0-10 (see the rule next to each axis) so the radar shape is a fair comparison. The same radar appears on the Dashboard."),
 ("Filtering", "Open Explorer and change Country, Status or Priority. The matching list, the selection KPIs and the status chart update instantly. Set a selector back to All to clear it."),
 ("Data quality rules", "Status, Priority, Level, Field and University are dropdown-only, Research Fit accepts 1-10 only. This keeps every chart and KPI reliable."),
 ("Before recording a demo", "Deadlines are fixed dates. If you record later, shift a few deadlines into the next 30 days so the Upcoming Deadlines list and the red alerts are visible."),
 ("Built with", "Python (openpyxl) generating native Excel formulas, validation, conditional formatting and chart objects. No macros, no external connections, works offline in Excel 2016 or later."),
]
hw["B3"] = "How It Works"; hw["B3"].font = f(16, True, NAVY)
for i, (k, v) in enumerate(rows, start=5):
    a = hw.cell(row=i, column=2, value=k); a.font = f(10, True, NAVY); a.alignment = Alignment(horizontal="left", vertical="top"); a.fill = fill(MID); a.border = BORDER
    b = hw.cell(row=i, column=3, value=v); b.font = f(10); b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True); b.border = BORDER
    hw.row_dimensions[i].height = 48
hw.sheet_view.showGridLines = False

# order + hide Lists
wb.move_sheet("Applications", offset=-(len(wb.sheetnames)-2))
order = ["Dashboard", "Applications", "Universities & Supervisors", "Explorer", "Scoring", "How It Works", "Lists"]
wb._sheets = [wb[n] for n in order]
lists.sheet_state = "hidden"
wb.active = 0
wb.save(OUT)
print("saved", OUT)
